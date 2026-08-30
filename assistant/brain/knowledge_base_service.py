import os
import glob
import threading
from logger import log, import_lock

class KnowledgeBaseService:
  def __init__(self, workspace_path="c:\\Users\\Sachin\\hackathon"):
    self.workspace_path = workspace_path
    self.docs_dir = os.path.join(workspace_path, "documents")
    self.db_dir = os.path.join(workspace_path, "brain", "chroma_db")
    self.lock = threading.Lock()
    self.ready = False
    
    # Ensure folders exist
    os.makedirs(self.docs_dir, exist_ok=True)
    
    self.vector_store = None
    self.embeddings = None
    
    # Initialize indexer in background thread to prevent startup latency
    threading.Thread(target=self._initialize_vector_db, daemon=True).start()

  def _initialize_vector_db(self):
    try:
      with import_lock:
        log.info("Initializing Local Knowledge Base (LangChain + ChromaDB + SentenceTransformers)...")
        from langchain_community.embeddings import HuggingFaceEmbeddings
        from langchain_community.vectorstores import Chroma
        
        try:
          # Try loading completely offline first (no network checks/latency)
          self.embeddings = HuggingFaceEmbeddings(
            model_name="all-MiniLM-L6-v2",
            model_kwargs={'device': 'cpu', 'local_files_only': True}
          )
          log.info("Hugging Face embeddings loaded from local cache.")
        except Exception as cache_err:
          log.warning(f"Could not load embeddings from local cache offline, trying online: {cache_err}")
          self.embeddings = HuggingFaceEmbeddings(
            model_name="all-MiniLM-L6-v2",
            model_kwargs={'device': 'cpu', 'local_files_only': False}
          )
        
        self.vector_store = Chroma(
          persist_directory=self.db_dir,
          embedding_function=self.embeddings
        )
      
      # Perform initial auto-indexing check
      self.index_documents()
      self.ready = True
      log.info("Local Knowledge Base initialized successfully!")
    except Exception as e:
      log.error(f"Failed to initialize vector database: {e}")

  def index_documents(self):
    """Scrapes files inside the documents directory and indexes them into the vector store."""
    with self.lock:
      if self.vector_store is None:
        log.warning("Vector store is not initialized. Skipping indexation.")
        return
        
      try:
        from langchain_text_splitters import RecursiveCharacterTextSplitter
        from langchain_core.documents import Document
        
        # List files
        files = []
        for ext in ["*.pdf", "*.docx", "*.pptx", "*.txt", "*.md"]:
          files.extend(glob.glob(os.path.join(self.docs_dir, ext)))
          
        if not files:
          log.info(f"No documents found to index in {self.docs_dir}")
          return
  
        new_docs = []
        for file_path in files:
          docs_from_file = self._parse_file(file_path)
          if docs_from_file:
            new_docs.extend(docs_from_file)
            
        if new_docs:
          # Split documents into chunks
          text_splitter = RecursiveCharacterTextSplitter(chunk_size=750, chunk_overlap=100)
          chunks = text_splitter.split_documents(new_docs)
          
          # Reset and recreate database to avoid stale or double indexed chunks
          # This keeps the database clean and lightweight
          from langchain_community.vectorstores import Chroma
          import shutil
          import gc
          
          # Release references and force garbage collection to close open SQLite/Chroma file locks on Windows
          self.vector_store = None
          gc.collect()
          
          try:
            if os.path.exists(self.db_dir):
              shutil.rmtree(self.db_dir)
          except Exception as se:
            log.error(f"Could not clear Chroma directory: {se}")
            
          self.vector_store = Chroma.from_documents(
            documents=chunks,
            embedding=self.embeddings,
            persist_directory=self.db_dir
          )
          log.info(f"Indexed {len(files)} files into {len(chunks)} chunks in ChromaDB.")
          
      except Exception as e:
        log.error(f"Failed during document indexing: {e}")

  def _parse_file(self, file_path: str) -> list:
    """Parses PDF, DOCX, PPTX, or TXT and returns a list of Document objects."""
    from langchain_core.documents import Document
    filename = os.path.basename(file_path)
    ext = os.path.splitext(filename)[1].lower()
    
    content = ""
    try:
      if ext == ".txt" or ext == ".md":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
          content = f.read()
          
      elif ext == ".pdf":
        import pypdf
        reader = pypdf.PdfReader(file_path)
        text_parts = []
        for page in reader.pages:
          text = page.extract_text()
          if text:
            text_parts.append(text)
        content = "\n".join(text_parts)
        
      elif ext == ".docx":
        import docx
        doc = docx.Document(file_path)
        text_parts = [p.text for p in doc.paragraphs]
        content = "\n".join(text_parts)
        
      elif ext == ".pptx":
        import pptx
        prs = pptx.Presentation(file_path)
        text_parts = []
        for slide in prs.slides:
          for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
              text_parts.append(shape.text)
        content = "\n".join(text_parts)
        
      if content.strip():
        return [Document(page_content=content, metadata={"source": filename, "path": file_path})]
    except Exception as e:
      log.error(f"Failed parsing file '{file_path}': {e}")
      
    return []

  def search_kb(self, query: str, k=4) -> str:
    """Searches the indexed documents and returns relevant chunks formatted as a context string."""
    if self.vector_store is None:
      return "Knowledge Base is still initializing, please wait!"
      
    try:
      results = self.vector_store.similarity_search(query, k=k)
      if not results:
        return "No relevant information found in your documents."
        
      formatted_parts = []
      for idx, doc in enumerate(results):
        source = doc.metadata.get("source", "Unknown")
        formatted_parts.append(f"[{idx+1}] Source: {source}\nContent: {doc.page_content}\n")
      return "\n---\n".join(formatted_parts)
    except Exception as e:
      log.error(f"Error querying knowledge base: {e}")
      return f"Error searching documents: {e}"
